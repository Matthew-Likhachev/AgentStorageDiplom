"""
import_module.py
Стратегия: Adapter Pattern, Data Transformation Pipeline
Модуль отвечает за преобразование внешнего формата (Excel) во внутреннее представление среды.
Объединяет загрузку, парсинг, валидацию и автоисправление карты.
"""

import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries


@dataclass
class ImportResult:
    """
    Стратегия: Result Object Pattern.
    Единый контракт возврата статуса импорта.
    """
    success: bool
    map_data: Optional[Dict[str, Any]] = None
    extracted_rules: Dict[str, Any] = field(default_factory=dict)  # Правила, извлеченные из карты
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    fixed_count: int = 0


class ExcelMapLoader:
    """Стратегия: Low-Level Data Adapter. Читает сырые данные из Excel."""

    VALID_TYPES = {'E', 'W', 'F', 'S', 'P', 'C', 'Z', 'B', 'A'}

    def __init__(self, file_path: Path):
        self.file_path = file_path
        self.map_data: Optional[Dict] = None
        self.errors: List[str] = []
        self.extracted_rules: Dict[str, Any] = {}

    def load(self) -> bool:
        """Загружает Excel и парсит в структуру данных"""
        if not self.file_path.exists():
            self.errors.append(f"Файл карты не найден: {self.file_path}")
            return False

        try:
            wb = load_workbook(self.file_path, data_only=True)
            ws = wb.active
        except Exception as e:
            self.errors.append(f"Ошибка открытия Excel: {e}")
            return False

        # === ОТЛАДКА: типы первых клеток ===
        logging.info("🔍 Отладка типов клеток первой строки:")
        for c in range(1, 6):
            cell = ws.cell(row=1, column=c)
            cell_text = str(cell.value) if cell.value else ''
            cell_type = cell_text.split('|')[0] if '|' in cell_text else 'N/A'
            logging.info(f"   {chr(64 + c)}1: тип='{cell_type}', текст='{cell_text[:30]}...'")
        # ============================================================

        logging.info("🔍 Поиск границ карты по типу 'E'...")
        width, height, start_row, start_col = self._find_bounds(ws)


        if width == 0 or height == 0:
            self.errors.append("Не удалось определить размеры карты (лист пуст или поврежден)")
            return False

        logging.info(f"✅ Карта найдена: начало=({start_row}, {start_col}), размер={width}x{height}")

        # Парсим ячейки
        self.map_data = {
            'width': width,
            'height': height,
            'start_row': start_row,
            'start_col': start_col,
            'cells': self._parse_grid(ws, start_row, start_col, width, height)
        }

        # === НОВОЕ: Извлекаем правила из содержимого карты ===
        self.extracted_rules = self._extract_rules_from_map()

        return True

    def _get_cell_color(self, cell) -> Optional[str]:
        """Безопасное получение HEX цвета ячейки (совместимо с openpyxl >= 3.0)"""
        fill = cell.fill
        if fill is None or fill.fgColor is None:
            return None

        rgb_val = fill.fgColor.rgb
        if rgb_val is None:
            return None

        # openpyxl может вернуть объект RGB или строку
        rgb_str = rgb_val.rgb if hasattr(rgb_val, 'rgb') else str(rgb_val)

        # openpyxl часто хранит цвета как 00RRGGBB (альфа-канал + RGB)
        if len(rgb_str) >= 8 and rgb_str.startswith('00'):
            return rgb_str[2:]
        elif len(rgb_str) == 6:
            return rgb_str

        return None

    def _get_cell_type(self, cell) -> str:
        """Извлекает тип клетки из текста (первая часть до '|')"""
        text = str(cell.value) if cell.value else ''
        return text.split('|')[0] if '|' in text else ''

    def _find_bounds(self, ws):
        """
        Стратегия: Auto-Detect Bounds by Cell Type.
        Определяет границы карты по типу клетки 'E' (Edge/стена).
        """
        start_r, start_c = 1, 1
        found_start = False

        # === Поиск первой клетки типа 'E' ===
        # Проверяем A1 в первую очередь
        first_cell = ws.cell(row=1, column=1)
        first_text = str(first_cell.value) if first_cell.value else ''
        first_type = first_text.split('|')[0] if '|' in first_text else ''

        logging.debug(f"Проверка A1: тип='{first_type}', текст='{first_text}'")

        if first_type == 'E':
            logging.info("✅ A1 имеет тип 'E'. Карта начинается с A1")
            start_r, start_c = 1, 1
            found_start = True
        else:
            # Ищем первую клетку типа 'E' в пределах 100x100
            logging.info(f"⚠️ A1 не типа 'E' (тип='{first_type}'). Ищем границу...")
            search_limit = 100
            max_r_search = min(ws.max_row or 1000, search_limit)
            max_c_search = min(ws.max_column or 1000, search_limit)

            for r in range(1, max_r_search + 1):
                for c in range(1, max_c_search + 1):
                    cell = ws.cell(row=r, column=c)
                    cell_text = str(cell.value) if cell.value else ''
                    cell_type = cell_text.split('|')[0] if '|' in cell_text else ''

                    if cell_type == 'E':
                        start_r, start_c = r, c
                        found_start = True
                        logging.info(f"✅ Найдена граница типа 'E' в ({r}, {c})")
                        break
                if found_start:
                    break

        # Fallback если не нашли границу типа 'E'
        if not found_start:
            logging.warning("⚠️ Граница типа 'E' не найдена. Использую границы данных листа.")
            if ws.dimensions:
                min_c, min_r, max_c, max_r = range_boundaries(ws.dimensions)
                return (max_c - min_c + 1), (max_r - min_r + 1), min_r, min_c
            else:
                return 0, 0, 1, 1

        # === Измеряем ширину (по типу 'E' в первой строке) ===
        width = 0
        for c in range(start_c, ws.max_column + 1):
            cell = ws.cell(row=start_r, column=c)
            cell_text = str(cell.value) if cell.value else ''
            cell_type = cell_text.split('|')[0] if '|' in cell_text else ''

            if cell_type == 'E':
                width += 1
            else:
                # Допускаем разрывы, если уже нашли несколько 'E'
                if width >= 2:
                    continue
                else:
                    break

        # Fallback для ширины
        if width == 0 and ws.dimensions:
            _, _, max_c, _ = range_boundaries(ws.dimensions)
            width = max_c - start_c + 1

        # === Измеряем высоту (по типу 'E' в первом столбце) ===
        height = 0
        for r in range(start_r, ws.max_row + 1):
            cell = ws.cell(row=r, column=start_c)
            cell_text = str(cell.value) if cell.value else ''
            cell_type = cell_text.split('|')[0] if '|' in cell_text else ''

            if cell_type == 'E':
                height += 1
            else:
                if height >= 2:
                    continue
                else:
                    break

        if height == 0 and ws.dimensions:
            _, max_r, _, _ = range_boundaries(ws.dimensions)
            height = max_r - start_r + 1

        logging.info(f"📐 Границы: начало=({start_r}, {start_c}), размер={width}x{height}")
        return width, height, start_r, start_c

    def _parse_grid(self, ws, start_row: int, start_col: int, width: int, height: int) -> List[List[Dict]]:
        """Парсит все ячейки карты с учетом найденных границ"""
        cells = []
        for r in range(start_row, start_row + height):
            row_data = []
            for c in range(start_col, start_col + width):
                cell = ws.cell(row=r, column=c)
                cell_text = str(cell.value) if cell.value else ''
                cell_type = self._get_cell_type(cell)

                row_data.append({
                    'row': r,
                    'col': c,
                    'text': cell_text,
                    'type': cell_type,  # Добавили тип
                    'color': self._get_cell_color(cell)
                })
            cells.append(row_data)
        return cells

    def _extract_rules_from_map(self) -> Dict[str, Any]:
        """
        Стратегия: Self-Describing Map.
        Автоматически извлекает конфигурацию системы (лимиты, слоты) из JSON-логики ячеек.
        """
        if not self.map_data:
            return {}

        packer_zones = {}    # {packer_id: [slots]}
        packer_buffers = {}  # {packer_id: [slots]}
        charger_ids = set()
        shelf_ids = set()

        # Проходим по всем ячейкам
        for row in self.map_data['cells']:
            for cell_info in row:
                text = cell_info.get('text', '')
                parts = text.split('|')
                if len(parts) < 4:
                    continue

                cell_type = parts[0]
                logic_str = parts[3]

                try:
                    logic = json.loads(logic_str) if logic_str.strip() and logic_str != '{}' else {}
                except json.JSONDecodeError:
                    continue

                pid = logic.get('packer_id')
                slot = logic.get('slot')

                # Сбор данных по типам
                if cell_type == 'C':
                    # Поддержка обоих ключей: charger_id и charger
                    cid = logic.get('charger_id') or logic.get('charger')
                    if cid is not None: charger_ids.add(cid)

                elif cell_type == 'S':
                    # Поддержка обоих ключей: shelf_id и shelf
                    sid = logic.get('shelf_id') or logic.get('shelf')
                    if sid is not None: shelf_ids.add(sid)

                elif cell_type == 'Z' and pid is not None and slot is not None:
                    # Зона заказа (обычно слоты 1-10)
                    packer_zones.setdefault(pid, []).append(slot)

                elif cell_type == 'B' and pid is not None and slot is not None:
                    # Буферная зона (обычно слоты 11+)
                    packer_buffers.setdefault(pid, []).append(slot)

        # Вычисляем лимиты
        # LOP лимит = максимальное количество слотов заказа у одного фасовщика
        max_packer_slots = 0
        max_buffer_slots = 0

        buffer_zones_map = {}

        for pid in packer_zones:
            z_slots = packer_zones[pid]
            b_slots = packer_buffers.get(pid, [])

            max_packer_slots = max(max_packer_slots, len(z_slots))
            max_buffer_slots = max(max_buffer_slots, len(b_slots))

            # Сохраняем маппинг для buffer_zones
            buffer_zones_map[f"packer_{pid}"] = len(b_slots)

        # Если ничего не нашли, ставим дефолт
        if max_packer_slots == 0: max_packer_slots = 10
        if max_buffer_slots == 0: max_buffer_slots = 10

        return {
            'order_rules': {
                'lop_limit': max_packer_slots,
                'buffer_slots': max_buffer_slots,
                'buffer_zones': buffer_zones_map
            },
            'lop_rules': {
                'packer_slots': max_packer_slots,
                'buffer_slots': max_buffer_slots
            },
            'system_info': {
                'packers': len(packer_zones),
                'chargers': len(charger_ids),
                'shelves': len(shelf_ids)
            }
        }


class MapValidator:
    """Валидатор данных карты (на основе рабочего оригинала)"""

    def __init__(self, map_data: Dict, auto_fix: bool = True):
        self.map_data = map_data
        self.auto_fix = auto_fix
        self.errors = []
        self.warnings = []
        self.fixed_count = 0

        self.width = map_data.get('width', 0)
        self.height = map_data.get('height', 0)
        self.start_row = map_data.get('start_row', 1)
        self.start_col = map_data.get('start_col', 1)
        self.cells = map_data.get('cells', [])

        self.COLOR_MAP = {
            '000000': 'E', 'FF0000': 'W', 'FFFFFF': 'F',
            'FF9900': 'S', '980000': 'P', '0000FF': 'C',
            '7030A0': 'Z', 'FFFF00': 'B', '00B050': 'A',
        }

        self.DIRECTIONS = ['n', 'w', 's', 'e']
        self.DIR_OFFSETS = {'n': (-1, 0), 's': (1, 0), 'w': (0, -1), 'e': (0, 1)}

    def parse_cell_text(self, cell_info: Dict) -> Tuple[Optional[Dict], Optional[str]]:
        """Парсит текст ячейки"""
        text = cell_info.get('text')
        row, col = cell_info['row'], cell_info['col']

        if text is None or not isinstance(text, str):
            return None, f"Ячейка ({row},{col}) пустая или содержит не текст"

        parts = text.split('|')
        if len(parts) != 4:
            return None, f"Ячейка ({row},{col}): неверный формат. Ожидалось 4 части, получено {len(parts)}"

        cell_type, loaded_dirs, empty_dirs, logic_str = parts

        if len(cell_type) != 1 or cell_type not in 'EFWSPCZBA':
            return None, f"Ячейка ({row},{col}): неверный тип '{cell_type}'"

        VALID_DIR_CHARS = set('nwse-')

        if len(loaded_dirs) > 7:
            return None, f"Ячейка ({row},{col}): неверная длина направлений с грузом '{loaded_dirs}'"
        if len(empty_dirs) > 7:
            return None, f"Ячейка ({row},{col}): неверная длина направлений без груза '{empty_dirs}'"
        if not all(c in VALID_DIR_CHARS for c in loaded_dirs):
            return None, f"Ячейка ({row},{col}): недопустимые символы в направлениях с грузом '{loaded_dirs}'"
        if not all(c in VALID_DIR_CHARS for c in empty_dirs):
            return None, f"Ячейка ({row},{col}): недопустимые символы в направлениях без груза '{empty_dirs}'"

        loaded = {'n': '-', 'w': '-', 's': '-', 'e': '-'}
        empty = {'n': '-', 'w': '-', 's': '-', 'e': '-'}

        for dir_char in self.DIRECTIONS:
            if dir_char in loaded_dirs:
                loaded[dir_char] = dir_char
            if dir_char in empty_dirs:
                empty[dir_char] = dir_char

        try:
            logic = json.loads(logic_str) if logic_str and logic_str.strip() and logic_str != '{}' else {}
        except json.JSONDecodeError:
            return None, f"Ячейка ({row},{col}): неверный JSON '{logic_str}'"

        return {
            'type': cell_type,
            'loaded_dirs': loaded,
            'empty_dirs': empty,
            'logic': logic,
            'row': row,
            'col': col,
            'original_text': text
        }, None

    def format_directions(self, dirs_dict: Dict) -> str:
        """Преобразует словарь направлений в строку формата n-w-s-e"""
        result = []
        for dir_char in self.DIRECTIONS:
            result.append(dirs_dict.get(dir_char, '-'))
            result.append('-')
        return ''.join(result[:-1])

    def update_cell_text(self, cell_info: Dict, cell_data: Dict) -> str:
        """Обновляет текст ячейки"""
        loaded_str = self.format_directions(cell_data['loaded_dirs'])
        empty_str = self.format_directions(cell_data['empty_dirs'])
        logic_str = json.dumps(cell_data['logic'], ensure_ascii=False) if cell_data['logic'] else '{}'
        new_text = f"{cell_data['type']}|{loaded_str}|{empty_str}|{logic_str}"

        r_idx = cell_info['row'] - self.start_row
        c_idx = cell_info['col'] - self.start_col
        if 0 <= r_idx < len(self.cells) and 0 <= c_idx < len(self.cells[0]):
            self.cells[r_idx][c_idx]['text'] = new_text
            self.fixed_count += 1
        return new_text

    def get_cell_info(self, row: int, col: int) -> Optional[Dict]:
        """Получает информацию о ячейке"""
        r_idx = row - self.start_row
        c_idx = col - self.start_col
        if 0 <= r_idx < self.height and 0 <= c_idx < self.width:
            return self.cells[r_idx][c_idx]
        return None

    def get_neighbor(self, row: int, col: int, direction: str) -> Optional[Dict]:
        """Возвращает соседнюю ячейку"""
        dr, dc = self.DIR_OFFSETS[direction]
        new_row, new_col = row + dr, col + dc
        cell_info = self.get_cell_info(new_row, new_col)
        if cell_info:
            cell_data, _ = self.parse_cell_text(cell_info)
            return cell_data
        return None

    def validate_cell_format(self) -> int:
        """Проверяет формат ячеек"""
        valid_count = 0
        for row_data in self.cells:
            for cell_info in row_data:
                cell_data, error = self.parse_cell_text(cell_info)
                if error:
                    self.errors.append(error)
                else:
                    valid_count += 1
                    color = cell_info.get('color')
                    expected_type = self.COLOR_MAP.get(color)
                    if expected_type and cell_data['type'] != expected_type:
                        self.warnings.append(
                            f"Ячейка ({cell_info['row']},{cell_info['col']}): тип '{cell_data['type']}' не соответствует цвету {color}")
        return valid_count

    def validate_movement_rules(self):
        """Проверяет и исправляет правила движения (ВАША ЛОГИКА)"""
        for row_data in self.cells:
            for cell_info in row_data:
                cell_data, _ = self.parse_cell_text(cell_info)
                if not cell_data:
                    continue

                row, col = cell_info['row'], cell_info['col']
                modified = False

                for direction in self.DIRECTIONS:
                    neighbor = self.get_neighbor(row, col, direction)
                    if not neighbor:
                        continue

                    neighbor_type = neighbor['type']

                    # 1. Нельзя в черные (E) и красные (W) — для любого агента
                    if neighbor_type in ['E', 'W']:
                        if cell_data['loaded_dirs'][direction] != '-':
                            if self.auto_fix:
                                cell_data['loaded_dirs'][direction] = '-'
                                modified = True
                            else:
                                self.errors.append(f"({row},{col}): движение с грузом в {neighbor_type}")
                        if cell_data['empty_dirs'][direction] != '-':
                            if self.auto_fix:
                                cell_data['empty_dirs'][direction] = '-'
                                modified = True
                            else:
                                self.errors.append(f"({row},{col}): движение без груза в {neighbor_type}")

                    # 2. Нельзя с грузом в полки (S) и зарядки (C)
                    elif neighbor_type in ['S', 'C']:
                        if cell_data['loaded_dirs'][direction] != '-':
                            if self.auto_fix:
                                cell_data['loaded_dirs'][direction] = '-'
                                modified = True
                            else:
                                self.errors.append(f"({row},{col}): движение с грузом в {neighbor_type}")

                    # 3. Нельзя в фасовщика (P) — для любого агента
                    elif neighbor_type == 'P':
                        if cell_data['loaded_dirs'][direction] != '-':
                            if self.auto_fix:
                                cell_data['loaded_dirs'][direction] = '-'
                                modified = True
                            else:
                                self.errors.append(f"({row},{col}): заезд с грузом в фасовщика")
                        if cell_data['empty_dirs'][direction] != '-':
                            if self.auto_fix:
                                cell_data['empty_dirs'][direction] = '-'
                                modified = True
                            else:
                                self.errors.append(f"({row},{col}): заезд без груза в фасовщика")

                    # 4. Нельзя без груза в зону заказа (Z, slot 1-10)
                    elif neighbor_type == 'Z':
                        slot = neighbor['logic'].get('slot', 0)
                        if 1 <= slot <= 10:
                            if cell_data['empty_dirs'][direction] != '-':
                                if self.auto_fix:
                                    cell_data['empty_dirs'][direction] = '-'
                                    modified = True
                                else:
                                    self.errors.append(f"({row},{col}): заезд без груза в зону заказа")

                    # === ВАЖНО: Если сосед — F, B или Z(slot>10), движение РАЗРЕШЕНО по умолчанию ===
                    # Ничего не делаем, просто пропускаем

                if modified and self.auto_fix:
                    self.update_cell_text(cell_info, cell_data)

    def validate_special_rules(self) -> Dict:
        """Проверяет уникальность ID packer/charger/shelf"""
        packer_ids, charger_ids, shelf_ids = set(), set(), set()

        for row_data in self.cells:
            for cell_info in row_data:
                cell_data, _ = self.parse_cell_text(cell_info)
                if not cell_data:
                    continue
                logic = cell_data['logic']
                cell_type = cell_data['type']
                row, col = cell_info['row'], cell_info['col']

                if cell_type == 'P' and 'packer' in logic:
                    pid = logic['packer']
                    if pid in packer_ids:
                        self.errors.append(f"Дублирующийся packer={pid}")
                    packer_ids.add(pid)
                elif cell_type == 'C' and 'charger' in logic:
                    cid = logic['charger']
                    if cid in charger_ids:
                        self.errors.append(f"Дублирующийся charger={cid}")
                    charger_ids.add(cid)
                elif cell_type == 'S' and 'shelf' in logic:
                    sid = logic['shelf']
                    if sid in shelf_ids:
                        self.errors.append(f"Дублирующийся shelf={sid}")
                    shelf_ids.add(sid)

        return {'packers': len(packer_ids), 'chargers': len(charger_ids), 'shelves': len(shelf_ids)}

    def validate(self) -> bool:
        self.validate_cell_format()
        self.validate_movement_rules()
        self.validate_special_rules()
        return len(self.errors) == 0


class MapImporter:
    """Стратегия: Facade / Pipeline Orchestrator."""

    def __init__(self, excel_path: str, auto_fix: bool = True, output_json: Optional[str] = None):
        self.loader = ExcelMapLoader(Path(excel_path))
        self.auto_fix = auto_fix
        self.output_json = Path(output_json) if output_json else None

    def import_map(self) -> ImportResult:
        # 1. Загрузка сырых данных
        if not self.loader.load():
            return ImportResult(success=False, errors=self.loader.errors)

        # 2. Валидация и автоисправление
        validator = MapValidator(self.loader.map_data, auto_fix=self.auto_fix)
        is_valid = validator.validate()

        result = ImportResult(
            success=is_valid,
            map_data=validator.map_data,
            extracted_rules=self.loader.extracted_rules, # Передаем правила
            errors=validator.errors,
            warnings=validator.warnings,
            fixed_count=validator.fixed_count
        )

        # 3. Сохранение очищенной версии (опционально)
        if self.output_json and validator.fixed_count > 0:
            try:
                self.output_json.parent.mkdir(parents=True, exist_ok=True)
                with open(self.output_json, 'w', encoding='utf-8') as f:
                    json.dump(validator.map_data, f, ensure_ascii=False, indent=2)
            except Exception as e:
                result.warnings.append(f"Не удалось сохранить фикс: {e}")

        return result