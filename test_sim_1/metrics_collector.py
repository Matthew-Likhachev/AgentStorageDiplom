"""
metrics_collector.py
Сбор метрик эффективности симуляции и сохранение в Excel.
"""
from __future__ import annotations
import os, time, datetime, logging
from dataclasses import dataclass
from typing import Dict, List, Optional, Set, Tuple


@dataclass
class _OrderRec:
    order_id: int; start_tick: int
    end_tick: Optional[int] = None; total_shelves: int = 0

@dataclass
class _SubRec:
    suborder_id: int; order_id: int; start_tick: int
    end_tick: Optional[int] = None; num_shelves: int = 0


class MetricsCollector:
    def __init__(self, algorithm_name: str, sim_params: dict):
        self.algorithm_name = algorithm_name
        self.sim_params     = sim_params
        self._wall_start_s  = time.time()
        self._wall_end_s    = 0.0
        self.sim_start_tick = 0
        self.sim_end_tick   = 0
        self._orders:    Dict[int, _OrderRec] = {}
        self._suborders: Dict[int, _SubRec]   = {}
        # Метрика 5
        self.waiting_events: int = 0
        self._blocked_this_tick: Set[int] = set()
        # Метрика 6
        self._queue_enter: Dict[int, Tuple[int, int]] = {}
        self._queue_wait_by_order: Dict[int, int] = {}
        # Метрики 7-9
        self.collision_count:       int = 0
        self.deadlock_count:        int = 0
        self.emergency_replan_count:int = 0
        self.current_tick: int = 0

    # ── Управление ───────────────────────────────────────────────────────────
    def on_sim_start(self, tick: int = 0):
        self.sim_start_tick = tick
        self._wall_start_s  = time.time()

    def on_sim_end(self, tick: int):
        self.sim_end_tick = tick
        self._wall_end_s  = time.time()

    def on_tick_start(self, tick: int):
        self.current_tick = tick
        self._blocked_this_tick.clear()

    # ── Заказы / Подзаказы ───────────────────────────────────────────────────
    def on_order_start(self, order_id: int, tick: int, total_shelves: int = 0):
        if order_id not in self._orders:
            self._orders[order_id] = _OrderRec(order_id, tick, total_shelves=total_shelves)

    def on_order_complete(self, order_id: int, tick: int):
        if order_id in self._orders and self._orders[order_id].end_tick is None:
            self._orders[order_id].end_tick = tick

    def on_suborder_start(self, suborder_id: int, order_id: int, tick: int, num_shelves: int = 0):
        if suborder_id not in self._suborders:
            self._suborders[suborder_id] = _SubRec(suborder_id, order_id, tick, num_shelves=num_shelves)

    def on_suborder_complete(self, suborder_id: int, tick: int):
        if suborder_id in self._suborders and self._suborders[suborder_id].end_tick is None:
            self._suborders[suborder_id].end_tick = tick

    # ── Метрики событий ──────────────────────────────────────────────────────
    def on_agent_blocked(self, agent_id: int, status: str):
        """Метрика 5: активный агент заблокирован Level-3 (не IDLE/WAITING_SLOT)."""
        if agent_id not in self._blocked_this_tick:
            self._blocked_this_tick.add(agent_id)
            self.waiting_events += 1

    def on_enter_packer_queue(self, agent_id: int, order_id: int, tick: int):
        """Метрика 6: агент встал в WAITING_SLOT."""
        self._queue_enter[agent_id] = (tick, order_id)

    def on_exit_packer_queue(self, agent_id: int, tick: int):
        """Метрика 6: агент покинул WAITING_SLOT."""
        if agent_id in self._queue_enter:
            enter_tick, order_id = self._queue_enter.pop(agent_id)
            self._queue_wait_by_order[order_id] = (
                self._queue_wait_by_order.get(order_id, 0) + tick - enter_tick)

    def on_collision(self):
        """Метрика 7: вынужденное ожидание при разрешении дедлока (Level-3)."""
        self.collision_count += 1

    def on_deadlock(self):
        """Метрика 8: Level-1 выдал REPLAN."""
        self.deadlock_count += 1

    def on_emergency_replan(self):
        """Метрика 9: Level-2 A* без мягких броней."""
        self.emergency_replan_count += 1

    # ── Вычисление сводки ────────────────────────────────────────────────────
    def compute_summary(self) -> dict:
        co  = [o for o in self._orders.values()    if o.end_tick is not None]
        cs  = [s for s in self._suborders.values() if s.end_tick is not None]
        n_o = len(co); n_s = len(cs)

        total_ticks  = self.sim_end_tick - self.sim_start_tick
        total_ms     = round((self._wall_end_s - self._wall_start_s) * 1000, 1)
        avg_ord_t    = round(sum(o.end_tick - o.start_tick for o in co) / n_o, 1) if n_o else 0
        avg_sh_ord   = round(sum(o.total_shelves for o in co) / n_o, 2)           if n_o else 0
        avg_sub_t    = round(sum(s.end_tick - s.start_tick for s in cs) / n_s, 1) if n_s else 0
        avg_sh_sub   = round(sum(s.num_shelves for s in cs) / n_s, 2)             if n_s else 0
        avg_queue    = round(sum(self._queue_wait_by_order.values()) / n_o, 1)    if n_o else 0

        return dict(
            algorithm=self.algorithm_name,
            num_agents=self.sim_params.get('num_agents', 0),
            num_orders=self.sim_params.get('num_orders', 0),
            order_seed=self.sim_params.get('order_seed', 0),
            shelf_seed=self.sim_params.get('shelf_seed', 0),
            total_ticks=total_ticks, total_ms=total_ms,
            completed_orders=n_o,
            avg_order_time_ticks=avg_ord_t, avg_shelves_per_order=avg_sh_ord,
            completed_suborders=n_s,
            avg_suborder_time_ticks=avg_sub_t, avg_shelves_per_suborder=avg_sh_sub,
            waiting_events=self.waiting_events,
            avg_queue_wait_per_order=avg_queue,
            collision_count=self.collision_count,
            deadlock_count=self.deadlock_count,
            emergency_replan_count=self.emergency_replan_count,
        )

    def save_to_excel_simple(self, output_dir: str, filename: str = 'metrics') -> str:
        """
        Сохраняет метрики в ПЛОСКИЙ Excel (1 заголовок + 1 строка данных).
        Предназначен для пакетного анализа через pandas.
        Имя файла: {filename}.xlsx
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            logging.error("[Metrics] openpyxl не установлен")
            return ''

        os.makedirs(output_dir, exist_ok=True)
        fpath = os.path.join(output_dir, f"{filename}.xlsx")

        s  = self.compute_summary()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "data"

        # Добавляем map_name если есть в sim_params
        row_data = {
            'map':                      self.sim_params.get('map_name', ''),
            'algorithm':                s['algorithm'],
            'num_agents':               s['num_agents'],
            'order_seed':               s['order_seed'],
            'shelf_seed':               s['shelf_seed'],
            'num_orders_target':        s['num_orders'],
            'total_ticks':              s['total_ticks'],
            'total_ms':                 s['total_ms'],
            'completed_orders':         s['completed_orders'],
            'avg_order_time_ticks':     s['avg_order_time_ticks'],
            'avg_shelves_per_order':    s['avg_shelves_per_order'],
            'completed_suborders':      s['completed_suborders'],
            'avg_suborder_time_ticks':  s['avg_suborder_time_ticks'],
            'avg_shelves_per_suborder': s['avg_shelves_per_suborder'],
            'waiting_events':           s['waiting_events'],
            'avg_queue_wait_per_order': s['avg_queue_wait_per_order'],
            'collision_count':          s['collision_count'],
            'deadlock_count':           s['deadlock_count'],
            'emergency_replan_count':   s['emergency_replan_count'],
        }

        hdr_font  = Font(name='Arial', bold=True, size=10, color='FFFFFFFF')
        hdr_fill  = PatternFill('solid', start_color='FF1F4E79', fgColor='FF1F4E79')
        hdr_align = Alignment(horizontal='center')

        for col, (key, val) in enumerate(row_data.items(), 1):
            h = ws.cell(row=1, column=col, value=key)
            h.font = hdr_font; h.fill = hdr_fill; h.alignment = hdr_align
            ws.cell(row=2, column=col, value=val)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max(14, len(key) + 2)

        wb.save(fpath)
        return fpath


        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            logging.error("[Metrics] openpyxl не установлен: pip install openpyxl")
            return ''

        os.makedirs(output_dir, exist_ok=True)
        ts    = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        fname = (f"{self.algorithm_name}"
                 f"_{self.sim_params.get('num_agents',0)}ag"
                 f"_{self.sim_params.get('num_orders',0)}ord"
                 f"_seed{self.sim_params.get('order_seed',0)}"
                 f"_{ts}.xlsx")
        fpath = os.path.join(output_dir, fname)

        s   = self.compute_summary()
        wb  = Workbook()
        ws  = wb.active
        ws.title = "Метрики"

        # --- Стили ---
        BLUE   = 'FF1F4E79'; WHITE = 'FFFFFFFF'
        LBLUE  = 'FFD6E4F0'; ALT   = 'FFF2F7FB'; WARN = 'FFFFF2CC'

        def _font(bold=False, sz=10, color='FF000000'):
            return Font(name='Arial', bold=bold, size=sz, color=color)

        def _fill(c): return PatternFill('solid', start_color=c, fgColor=c)

        thin   = Side(style='thin', color='FFB0C4DE')
        bdr    = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left   = Alignment(horizontal='left',   vertical='center', wrap_text=True)

        ws.column_dimensions['A'].width =  5
        ws.column_dimensions['B'].width = 44
        ws.column_dimensions['C'].width = 22
        ws.column_dimensions['D'].width = 42

        # Заголовок
        ws.merge_cells('A1:D1')
        c = ws['A1']
        c.value = f'Метрики симуляции — {s["algorithm"].upper()}'
        c.font  = Font(name='Arial', bold=True, size=14, color=WHITE)
        c.fill  = _fill(BLUE); c.alignment = center
        ws.row_dimensions[1].height = 28

        # Параметры строкой
        ws.merge_cells('A2:D2')
        c = ws['A2']
        c.value = (f"Агентов: {s['num_agents']}  |  Заказов: {s['num_orders']}  |  "
                   f"order_seed: {s['order_seed']}  |  shelf_seed: {s['shelf_seed']}  |  "
                   f"Дата: {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}")
        c.font  = Font(name='Arial', italic=True, size=9, color=BLUE)
        c.fill  = _fill(LBLUE); c.alignment = center
        ws.row_dimensions[2].height = 18

        row = 4

        def section(r, title):
            ws.merge_cells(f'A{r}:D{r}')
            c = ws[f'A{r}']
            c.value = f'  {title}'
            c.font  = Font(name='Arial', bold=True, size=10, color=BLUE)
            c.fill  = _fill(LBLUE); c.alignment = left
            ws.row_dimensions[r].height = 20
            return r + 1

        def mrow(r, num, label, value, note='', warn=False):
            bg = WARN if warn else (ALT if r % 2 == 0 else WHITE)
            for col in range(1, 5):
                ce = ws.cell(row=r, column=col)
                ce.fill = _fill(bg); ce.border = bdr
            ws.cell(r, 1).value = num
            ws.cell(r, 1).font  = _font(bold=True, color=BLUE)
            ws.cell(r, 1).alignment = center
            ws.cell(r, 2).value = label
            ws.cell(r, 2).font  = _font(); ws.cell(r, 2).alignment = left
            ws.cell(r, 3).value = value
            ws.cell(r, 3).font  = _font(bold=True); ws.cell(r, 3).alignment = center
            ws.cell(r, 4).value = note
            ws.cell(r, 4).font  = Font(name='Arial', size=8, italic=True, color='FF606060')
            ws.cell(r, 4).alignment = left
            ws.row_dimensions[r].height = 18
            return r + 1

        row = section(row, '⚙  Параметры')
        row = mrow(row, '—', 'Алгоритм',              s['algorithm'])
        row = mrow(row, '—', 'Количество агентов',    s['num_agents'])
        row = mrow(row, '—', 'Заказов в сценарии',    s['num_orders'])
        row = mrow(row, '—', 'Seed заказов / стеллажей',
                              f"{s['order_seed']} / {s['shelf_seed']}")
        row += 1

        row = section(row, '⏱  Время выполнения сценария')
        row = mrow(row, '1', 'Общее время — тактов',
                              f"{s['total_ticks']:,}",
                              'От первого тика до завершения последнего заказа')
        row = mrow(row, '2', 'Общее время — мс (реальное)',
                              f"{s['total_ms']:,.1f} мс",
                              'Wall-clock время работы Python-процесса')
        row += 1

        row = section(row, '📦  Заказы и подзаказы')
        row = mrow(row, '—', 'Завершено заказов',        s['completed_orders'])
        row = mrow(row, '3', 'Среднее время заказа — тактов',
                              f"{s['avg_order_time_ticks']:.1f}",
                              f"Ср. стеллажей в заказе: {s['avg_shelves_per_order']:.2f}")
        row = mrow(row, '—', 'Завершено подзаказов',     s['completed_suborders'])
        row = mrow(row, '4', 'Среднее время подзаказа — тактов',
                              f"{s['avg_suborder_time_ticks']:.1f}",
                              f"Ср. стеллажей в подзаказе: {s['avg_shelves_per_suborder']:.2f}")
        row += 1

        row = section(row, '🤖  Поведение агентов')
        row = mrow(row, '5', 'Ситуации ожидания (активные агенты заблокированы)',
                              f"{s['waiting_events']:,}",
                              'GO/CARRY/RETURN стоят из-за других. IDLE и WAITING_SLOT не считаются.',
                              warn=s['waiting_events'] > 1000)
        row = mrow(row, '6', 'Среднее ожидание в очереди у фасовщика — тактов/заказ',
                              f"{s['avg_queue_wait_per_order']:.1f}",
                              'Суммарные тики в WAITING_SLOT / число завершённых заказов')
        row += 1

        row = section(row, '⚠  Конфликты')
        row = mrow(row, '7', 'Столкновений (вынужденных ожиданий при дедлоке)',
                              f"{s['collision_count']:,}",
                              'Swap и vertex force_wait из Level-3',
                              warn=s['collision_count'] > 300)
        row = mrow(row, '8', 'Дедлоков (REPLAN от Level-1)',
                              f"{s['deadlock_count']:,}",
                              'Агент не двигался > stuck_threshold тиков',
                              warn=s['deadlock_count'] > 100)
        row = mrow(row, '9', 'Экстренных перепланирований маршрутов',
                              f"{s['emergency_replan_count']:,}",
                              'A* нашёл путь только без временны́х броней',
                              warn=s['emergency_replan_count'] > 200)

        # ── Лист 2: Заказы ────────────────────────────────────────────────────
        ws2 = wb.create_sheet("Заказы")
        hdrs = ['order_id','start_tick','end_tick','duration_ticks',
                'total_shelves','queue_wait_ticks']
        for ci, h in enumerate(hdrs, 1):
            c = ws2.cell(1, ci, h)
            c.font = Font(name='Arial', bold=True, size=10, color=WHITE)
            c.fill = _fill(BLUE); c.alignment = center
        for ri, o in enumerate(sorted(self._orders.values(), key=lambda x: x.order_id), 2):
            dur   = (o.end_tick - o.start_tick) if o.end_tick else None
            qw    = self._queue_wait_by_order.get(o.order_id, 0)
            for ci, v in enumerate([o.order_id,o.start_tick,o.end_tick,dur,o.total_shelves,qw],1):
                ws2.cell(ri, ci, v).font = _font()
        for ci in range(1, len(hdrs)+1):
            ws2.column_dimensions[get_column_letter(ci)].width = 18

        # ── Лист 3: Подзаказы ────────────────────────────────────────────────
        ws3 = wb.create_sheet("Подзаказы")
        hdrs3 = ['suborder_id','order_id','start_tick','end_tick',
                 'duration_ticks','num_shelves']
        for ci, h in enumerate(hdrs3, 1):
            c = ws3.cell(1, ci, h)
            c.font = Font(name='Arial', bold=True, size=10, color=WHITE)
            c.fill = _fill(BLUE); c.alignment = center
        for ri, su in enumerate(sorted(self._suborders.values(), key=lambda x: x.suborder_id), 2):
            dur_ = (su.end_tick - su.start_tick) if su.end_tick else None
            for ci, v in enumerate([su.suborder_id,su.order_id,su.start_tick,
                                     su.end_tick,dur_,su.num_shelves], 1):
                ws3.cell(ri, ci, v).font = _font()
        for ci in range(1, len(hdrs3)+1):
            ws3.column_dimensions[get_column_letter(ci)].width = 16

        wb.save(fpath)
        logging.info(f"[MetricsCollector] Сохранено: {fpath}")
        return fpath